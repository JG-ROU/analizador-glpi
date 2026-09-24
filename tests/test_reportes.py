"""Series, gráficos y reportes REP-01 y REP-03 en Excel y CSV."""

from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from ayudas import importar_diario
from core import seguridad
from core.analisis import series
from core.analisis.filtros import Filtros
from core.analisis.kpis import CalculadoraKPI
from core.errores import ErrorPermiso, ErrorValidacion
from core.reportes import catalogo, graficos
from core.reportes.formatos import anonimizar
from test_kpis import AHORA, SEPTIEMBRE, escenario, tecnico_id


def d(mes, dia):
    return datetime(2026, mes, dia, 10, 0)


@pytest.fixture
def datos(bd, coordinador, tmp_path):
    importar_diario(bd, coordinador, escenario(), d(8, 25), datetime(2026, 10, 3, 23, 0), tmp_path)
    with bd:
        bd.execute("UPDATE parametro SET valor = '72' WHERE clave = 'sla_horas_urgente'")
        bd.execute("UPDATE parametro SET valor = '96' WHERE clave = 'sla_horas_mediana'")
    return bd


def solicitud(bd, sesion, tmp_path, **otros):
    return catalogo.SolicitudReporte(
        conexion=bd, sesion=sesion, periodo=SEPTIEMBRE,
        carpeta_exportaciones=tmp_path / "exportaciones", ahora=AHORA, **otros,
    )


# --- Series ---

def test_recibidos_resueltos_por_semana(datos, coordinador):
    serie = series.recibidos_resueltos(CalculadoraKPI(datos, coordinador, AHORA), SEPTIEMBRE)
    assert [p.etiqueta for p in serie] == ["S36", "S37", "S38", "S39", "S40"]
    assert sum(p.valores["Recibidos"] for p in serie) == 10
    assert serie[0].valores == {"Recibidos": 4.0, "Resueltos": 3.0}


def test_periodos_largos_se_agrupan_por_mes(datos, coordinador):
    from core.analisis import periodos
    anual = periodos.rango(datetime(2026, 1, 1).date(), datetime(2026, 12, 31).date())
    assert len(series.subperiodos(anual)) == 12


def test_carga_por_tecnico(datos, coordinador):
    carga = series.carga_por_tecnico(datos, CalculadoraKPI(datos, coordinador, AHORA), SEPTIEMBRE)
    assert carga == {"Tecnico 03": {"Mediana": 1}, "Tecnico 04": {"Mediana": 6}}


def test_carga_por_tecnico_para_consulta_solo_la_propia(datos, coordinador, consulta):
    usuario = seguridad.crear_usuario(
        datos, coordinador, nombre="T4", perfil=seguridad.CONSULTA, pin="4444",
        tecnico_id=tecnico_id(datos, "Tecnico 04"),
    )
    sesion = seguridad.iniciar_sesion(datos, usuario, "4444")
    carga = series.carga_por_tecnico(datos, CalculadoraKPI(datos, sesion, AHORA), SEPTIEMBRE)
    assert list(carga) == ["Tecnico 04"]
    assert series.carga_por_tecnico(datos, CalculadoraKPI(datos, consulta, AHORA), SEPTIEMBRE) == {}


def test_recibidos_por_estado(datos, coordinador):
    conteo = series.recibidos_por_estado(datos, CalculadoraKPI(datos, coordinador, AHORA), SEPTIEMBRE)
    assert conteo["En curso (asignado)"] == 6 and conteo["Resuelto"] == 3 and conteo["Cerrado"] == 1


def test_brechas(datos, coordinador):
    calc = CalculadoraKPI(datos, coordinador, AHORA)
    brechas = series.brechas(calc.calcular_visibles(SEPTIEMBRE), calc.definiciones)
    kpi04 = next(b for b in brechas if b.codigo == "KPI-04")
    assert (kpi04.valor, kpi04.objetivo, kpi04.tipo_objetivo) == (80.0, 100, "Meta")
    kpi05 = next(b for b in brechas if b.codigo == "KPI-05")
    assert (kpi05.objetivo, kpi05.tipo_objetivo) == (20, "Umbral verde")


# --- Gráficos ---

def test_graficos_generan_png(datos, coordinador):
    calc = CalculadoraKPI(datos, coordinador, AHORA)
    figuras = [
        graficos.recibidos_resueltos(series.recibidos_resueltos(calc, SEPTIEMBRE)),
        graficos.backlog(series.backlog(calc, SEPTIEMBRE)),
        graficos.barras(series.abiertos_por_prioridad(calc, SEPTIEMBRE), "Abiertos", graficos.COLORES_PRIORIDAD),
        graficos.carga_por_tecnico(series.carga_por_tecnico(datos, calc, SEPTIEMBRE)),
        graficos.brecha_vs_meta(series.brechas(calc.calcular_visibles(SEPTIEMBRE), calc.definiciones)),
    ]
    for figura in figuras:
        assert graficos.a_png(figura).startswith(b"\x89PNG")


def test_graficos_sin_datos():
    for figura in (graficos.recibidos_resueltos([]), graficos.backlog([]), graficos.barras({"a": 0}, "x"),
                   graficos.carga_por_tecnico({}), graficos.brecha_vs_meta([])):
        assert graficos.a_png(figura).startswith(b"\x89PNG")


# --- REP-01 ---

def test_rep01_excel(datos, coordinador, tmp_path):
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path), "REP-01", catalogo.EXCEL)
    assert ruta.parent.name == "2026-10"
    assert ruta.name.startswith("REP-01_2026-09_20261005_120000")
    hoja = load_workbook(ruta)["Resumen"]
    celdas = {c.value for fila in hoja.iter_rows() for c in fila if c.value is not None}
    assert "REP-01 Resumen de indicadores" in celdas
    assert "Septiembre 2026" in celdas and "Sin filtros" in celdas and "Coordinador" in celdas
    assert "80 %" in celdas and "✔ Verde" in celdas and "✖ Rojo" in celdas
    assert any(isinstance(v, str) and v.startswith("Criticidad global del período: ✖ Rojo") for v in celdas)
    assert len(hoja._images) == 2
    fila_kpi04 = next(f for f in hoja.iter_rows() if f[0].value == "KPI-04")
    semaforo_celda = fila_kpi04[4]
    assert semaforo_celda.value == "✔ Verde" and semaforo_celda.fill.fgColor.rgb.endswith("C8E6C9")


def test_rep01_csv(datos, coordinador, tmp_path):
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path), "REP-01", catalogo.CSV)
    assert ruta.read_bytes().startswith(b"\xef\xbb\xbf")
    tabla = pd.read_csv(ruta, sep=";", encoding="utf-8-sig", dtype=str)
    assert list(tabla.columns)[:5] == ["Código", "Indicador", "Valor", "Meta", "Semáforo"]
    fila = tabla[tabla["Código"] == "KPI-05"].iloc[0]
    assert (fila["Valor"], fila["Semáforo"]) == ("50 %", "✖ Rojo")


def test_rep01_con_filtros_descritos(datos, coordinador, tmp_path):
    filtros = Filtros(tecnico_id=tecnico_id(datos, "Tecnico 03"), turno="Mañana")
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path, filtros=filtros), "REP-01", catalogo.EXCEL)
    celdas = {c.value for f in load_workbook(ruta)["Resumen"].iter_rows() for c in f}
    assert "Técnico: Tecnico 03; Turno: Mañana" in celdas


# --- REP-03 ---

def test_rep03_coordinador(datos, coordinador, tmp_path):
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path), "REP-03", catalogo.CSV)
    tabla = pd.read_csv(ruta, sep=";", encoding="utf-8-sig")
    assert list(tabla["Técnico"]) == ["Tecnico 01", "Tecnico 02", "Tecnico 03", "Tecnico 04"]
    t2 = tabla[tabla["Técnico"] == "Tecnico 02"].iloc[0]
    assert (t2["Atendidos"], t2["% SLA"], t2["Resueltos sin cerrar (informativo)"]) == (2, 100.0, 1)
    excel = catalogo.generar(solicitud(datos, coordinador, tmp_path), "REP-03", catalogo.EXCEL)
    celdas = {c.value for f in load_workbook(excel)["Responsables"].iter_rows() for c in f}
    assert "El número de tickets es informativo (carga), no una calificación del técnico." in celdas


def test_rep03_consulta_solo_sus_metricas(datos, coordinador, tmp_path):
    usuario = seguridad.crear_usuario(
        datos, coordinador, nombre="T2", perfil=seguridad.CONSULTA, pin="2222",
        tecnico_id=tecnico_id(datos, "Tecnico 02"),
    )
    sesion = seguridad.iniciar_sesion(datos, usuario, "2222")
    ruta = catalogo.generar(solicitud(datos, sesion, tmp_path), "REP-03", catalogo.CSV)
    assert list(pd.read_csv(ruta, sep=";", encoding="utf-8-sig")["Técnico"]) == ["Tecnico 02"]
    otro = Filtros(tecnico_id=tecnico_id(datos, "Tecnico 01"))
    with pytest.raises(ErrorPermiso):
        catalogo.generar(solicitud(datos, sesion, tmp_path, filtros=otro), "REP-03", catalogo.CSV)


def test_reporte_o_formato_inexistente(datos, coordinador, tmp_path):
    with pytest.raises(ErrorValidacion):
        catalogo.generar(solicitud(datos, coordinador, tmp_path), "REP-99", catalogo.EXCEL)
    with pytest.raises(ErrorValidacion):
        catalogo.generar(solicitud(datos, coordinador, tmp_path), "REP-01", "WORD")


# --- PDF (CA-13) y reportes de la Fase 2 ---

def texto_pdf(ruta):
    from pypdf import PdfReader
    lector = PdfReader(ruta)
    return lector, "\n".join(pagina.extract_text() for pagina in lector.pages)


def test_ca13_pdf_rep01_con_filtros_fecha_usuario_graficos_y_paginacion(datos, coordinador, tmp_path):
    filtros = Filtros(turno="Mañana")
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path, filtros=filtros), "REP-01", catalogo.PDF)
    lector, texto = texto_pdf(ruta)
    assert "REP-01 Resumen de indicadores" in texto
    assert "Filtros: Turno: Mañana" in texto
    assert "Generado: 05/10/2026 12:00" in texto and "Usuario: Coordinador" in texto
    assert f"Página 1 de {len(lector.pages)}" in texto
    assert "✔ Verde" in texto and "✖ Rojo" in texto
    imagenes = sum(len(pagina.images) for pagina in lector.pages)
    assert imagenes == 2


@pytest.mark.parametrize("codigo", ["REP-02", "REP-03", "REP-04", "REP-05", "REP-09", "REP-11"])
@pytest.mark.parametrize("formato", [catalogo.PDF, catalogo.EXCEL, catalogo.CSV])
def test_reportes_de_fase_2_en_los_tres_formatos(datos, coordinador, tmp_path, codigo, formato):
    from core import clasificacion as cl
    from core.analisis import hallazgos
    norte = cl.guardar_estacion(datos, coordinador, nombre="Norte", cliente="Cliente A")
    cl.clasificar(datos, coordinador, [11, 12, 13, 14], {cl.ESTACION: norte, cl.CATEGORIA: "REC-01"})
    hallazgos.detectar(datos, ahora=AHORA)
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path), codigo, formato)
    assert ruta.exists() and ruta.stat().st_size > 0
    if formato == catalogo.PDF:
        _, texto = texto_pdf(ruta)
        assert codigo in texto


def test_pdf_con_celdas_enormes_no_falla(tmp_path):
    from core.reportes.formatos import Hoja, Reporte, Tabla, escribir_pdf
    tickets = ", ".join(str(n) for n in range(100000, 103000))
    tabla = Tabla("Hallazgos", pd.DataFrame({"Regla": ["HAL-05"] * 50, "Tickets": [tickets] * 50}))
    ruta = escribir_pdf(Reporte("Prueba", {"Usuario": "X"}, [Hoja("H", [tabla])]), tmp_path / "grande.pdf")
    _, texto = texto_pdf(ruta)
    assert "ver el Excel" in " ".join(texto.split())


def test_rep08_segmov_guarda_la_distribucion(datos, coordinador, tmp_path):
    from core import clasificacion as cl
    norte = cl.guardar_estacion(datos, coordinador, nombre="Norte")
    sur = cl.guardar_estacion(datos, coordinador, nombre="Sur")
    cl.clasificar(datos, coordinador, [11, 12, 13], {cl.ESTACION: norte})
    cl.clasificar(datos, coordinador, [14], {cl.ESTACION: sur})
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path, total_horas=100), "REP-08", catalogo.CSV)
    tabla = pd.read_csv(ruta, sep=";", encoding="utf-8-sig")
    assert dict(zip(tabla["Estación"], tabla["Horas"])) == {"Norte": 75.0, "Sur": 25.0}


def test_reportes_solo_coordinador(datos, consulta, tmp_path):
    for codigo in ("REP-08", "REP-11"):
        with pytest.raises(ErrorPermiso):
            catalogo.generar(solicitud(datos, consulta, tmp_path, total_horas=100), codigo, catalogo.PDF)
    assert {r.codigo for r in catalogo.disponibles(consulta)} == {"REP-01", "REP-02", "REP-03", "REP-04", "REP-05", "REP-09"}


def test_rep11_anonimizado(datos, coordinador, tmp_path):
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path, anonimizar=True), "REP-11", catalogo.CSV)
    tabla = pd.read_csv(ruta, sep=";", encoding="utf-8-sig", dtype=str)
    assert tabla["Técnico"].str.startswith("Persona").all() and tabla["Autor"].str.startswith("Persona").all()
    assert len(tabla) == 10


def test_rep09_fuera_de_sla(datos, coordinador, tmp_path):
    ruta = catalogo.generar(solicitud(datos, coordinador, tmp_path), "REP-09", catalogo.EXCEL)
    celdas = {c.value for f in load_workbook(ruta)["SLA"].iter_rows() for c in f}
    assert "Tickets fuera de SLA o en riesgo" in celdas and "✖ Incumplido" in celdas


# --- Anonimización (IMP-05, RNF-08) ---

def test_anonimizar_autores():
    datos = pd.DataFrame({"Autor": ["Ana", "Luis", "Ana", "", None], "Título": ["a", "b", "c", "d", "e"]})
    resultado = anonimizar(datos, ("Autor", "No existe"))
    assert list(resultado["Autor"][:4]) == ["Persona 001", "Persona 002", "Persona 001", ""]
    assert pd.isna(resultado["Autor"][4])
    assert list(datos["Autor"][:2]) == ["Ana", "Luis"]  # el original no cambia
