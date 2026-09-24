"""Notificaciones (NOT-01, 02, 03, 05), paquete mensual y borradores .eml (CA-14)."""

from datetime import datetime
from email import policy
from email.parser import BytesParser

import pytest

import generador_glpi as gen
from ayudas import importar_diario, ticket
from core import configuracion as cfg
from core import notificaciones as noti
from core.analisis import hallazgos, periodos
from core.errores import ErrorPermiso, ErrorValidacion
from core.reportes import paquete


def d(mes, dia, hora=10):
    return datetime(2026, mes, dia, hora, 0)


@pytest.fixture
def datos(bd, coordinador, tmp_path):
    tickets = [
        ticket(1, d(8, 3), "Tecnico 01", "Mayor", (d(8, 4), gen.RESUELTO)),
        ticket(2, d(8, 5), "Tecnico 01", "Baja", (d(8, 6), gen.ESCALADO)),
        ticket(3, d(9, 2), "Tecnico 02", "Mayor"),
    ]
    importar_diario(bd, coordinador, tickets, d(8, 3), d(9, 2, 23), tmp_path)
    return bd


def lee_eml(ruta):
    with ruta.open("rb") as archivo:
        return BytesParser(policy=policy.default).parse(archivo)


def codigos(lista):
    return [n.codigo for n in lista]


def test_ca14_not03_al_abrir_en_mes_nuevo_y_paquete_mensual(datos, coordinador, tmp_path):
    ahora = d(9, 3, 8)
    assert "NOT-03" in codigos(noti.pendientes(datos, coordinador, noti.AL_INICIAR, ahora))
    cfg.actualizar_parametro(datos, coordinador, "correo_jefatura", "jefatura@empresa.test")
    agosto = periodos.mes(2026, 8)
    resultado = paquete.generar(datos, coordinador, agosto, tmp_path / "exportaciones",
                                resumen_ejecutivo="Mes estable.\nSe reforzó el turno nocturno.",
                                plan_mejora="Reducir escalamientos.", ahora=ahora)
    assert resultado.pdf.exists() and resultado.excel.exists() and resultado.correo.exists()
    mensaje = lee_eml(resultado.correo)
    assert mensaje["X-Unsent"] == "1" and mensaje["To"] == "jefatura@empresa.test"
    assert "Agosto 2026" in mensaje["Subject"]
    adjuntos = [parte.get_filename() for parte in mensaje.iter_attachments()]
    assert adjuntos == [resultado.pdf.name]
    assert "Mes estable." in mensaje.get_body(preferencelist=("plain",)).get_content()
    assert "NOT-03" not in codigos(noti.pendientes(datos, coordinador, noti.AL_INICIAR, ahora))


def test_paquete_incluye_las_secciones(datos, coordinador, tmp_path):
    from openpyxl import load_workbook
    from pypdf import PdfReader
    resultado = paquete.generar(datos, coordinador, periodos.mes(2026, 8), tmp_path, ahora=d(9, 3))
    hojas = load_workbook(resultado.excel).sheetnames
    assert hojas == ["Resumen ejecutivo", "Resumen", "SLA", "Tendencias", "Hallazgos", "SEGMOV",
                     "Calidad del área", "Plan de mejora"]
    texto = "\n".join(p.extract_text() for p in PdfReader(resultado.pdf).pages)
    assert "REP-10 Paquete mensual gerencial" in texto and "SEGMOV del mes no generado" in texto


def test_paquete_solo_coordinador_y_por_mes(datos, consulta, coordinador, tmp_path):
    with pytest.raises(ErrorPermiso):
        paquete.generar(datos, consulta, periodos.mes(2026, 8), tmp_path)
    with pytest.raises(ErrorValidacion, match="un mes"):
        paquete.generar(datos, coordinador, periodos.semana(2026, 36), tmp_path)


def test_not01_importacion_desactualizada(bd, coordinador, datos):
    assert "NOT-01" in codigos(noti.pendientes(bd, coordinador, noti.AL_INICIAR, d(12, 1)))


def test_not02_y_not05_despues_de_importar(datos, coordinador, tmp_path):
    hallazgos.detectar(datos, ahora=d(9, 3))
    despues = noti.pendientes(datos, coordinador, noti.DESPUES_DE_IMPORTAR, d(9, 3))
    assert "NOT-02" in codigos(despues)  # P1 abierto = ALTA


def test_not05_kpi_critico_en_rojo(bd, coordinador, tmp_path):
    tickets = [ticket(1, d(9, 1), "Tecnico 01", "Baja", (d(9, 2), gen.ESCALADO))]
    importar_diario(bd, coordinador, tickets, d(9, 1), d(9, 2, 23), tmp_path)
    despues = noti.pendientes(bd, coordinador, noti.DESPUES_DE_IMPORTAR, d(9, 3))
    rojo = next(n for n in despues if n.codigo == "NOT-05")
    assert "KPI-05" in rojo.mensaje


def test_consulta_no_recibe_not03(datos, consulta):
    assert "NOT-03" not in codigos(noti.pendientes(datos, consulta, noti.AL_INICIAR, d(9, 3)))


def test_borrador_de_hallazgos_altos(datos, coordinador, tmp_path):
    with pytest.raises(ErrorValidacion, match="No hay hallazgos"):
        paquete.borrador_hallazgos_altos(datos, coordinador, tmp_path, ahora=d(9, 3))
    hallazgos.detectar(datos, ahora=d(9, 3))
    ruta = paquete.borrador_hallazgos_altos(datos, coordinador, tmp_path, ahora=d(9, 3))
    assert "P1 abierto" in lee_eml(ruta).get_body(preferencelist=("plain",)).get_content()


def test_parametro_de_texto(bd, coordinador):
    cfg.actualizar_parametro(bd, coordinador, "correo_jefatura", "  a@b.test,  c@d.test ")
    assert bd.execute("SELECT valor FROM parametro WHERE clave = 'correo_jefatura'").fetchone()[0] == "a@b.test, c@d.test"
    cfg.actualizar_parametro(bd, coordinador, "correo_jefatura", "")
    assert bd.execute("SELECT valor FROM parametro WHERE clave = 'correo_jefatura'").fetchone()[0] is None
