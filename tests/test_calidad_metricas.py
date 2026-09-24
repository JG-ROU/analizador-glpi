"""Métricas de calidad: CAL-04 a CAL-08 (CA-19 a CA-22) y NOT-04."""

from datetime import datetime, timedelta

import pytest

import generador_glpi as gen
from ayudas import importar_diario, ticket
from core import clasificacion as cl
from core import notificaciones as noti
from core import seguridad
from core.analisis import calidad as cal
from core.analisis import calidad_metricas as met
from core.analisis import periodos
from core.errores import ErrorPermiso


def d(mes, dia, hora=10, minuto=0):
    return datetime(2026, mes, dia, hora, minuto)


def tecnico_id(bd, nombre):
    return bd.execute("SELECT id FROM tecnico WHERE nombre_glpi = ?", (nombre,)).fetchone()[0]


def evaluar_todo(bd, sesion, ticket_id, valor=cal.CUMPLE, ahora=None, fallar=(), retroalimentacion=""):
    tipo = cal.tipo_caso_sugerido(bd, ticket_id)
    automaticos = cal.precalificar(bd, ticket_id, tipo, ahora)
    resultados = {}
    for c in cal.criterios_para(bd, tipo):
        deseado = cal.NO_CUMPLE if c["id"] in fallar else valor
        nota = "ajuste de prueba" if c["id"] in automaticos and automaticos[c["id"]].resultado != deseado else None
        resultados[c["id"]] = (deseado, nota)
    return cal.guardar_evaluacion(bd, sesion, ticket_id, tipo, resultados, retroalimentacion, ahora=ahora)


# --- CAL-04 (CA-19) ---

@pytest.fixture
def semana_cerrada(bd, coordinador, tmp_path):
    """Semana 36 (31-ago a 6-sep): 1 P1, 1 escalado largo y tickets de 4 técnicos."""
    tickets = [
        ticket(1, d(9, 1), "Tecnico 01", "Mayor", (d(9, 1, 14), gen.RESUELTO), (d(9, 2), gen.CERRADO)),
        ticket(2, d(8, 25), "Tecnico 02", "Baja", (d(8, 26), gen.ESCALADO), (d(9, 3), gen.RESUELTO), (d(9, 4), gen.CERRADO)),
    ]
    for n in range(12):
        tickets.append(ticket(10 + n, d(9, 1, 8 + n), f"Tecnico 0{3 + n % 2}", "Baja",
                              (d(9, 2, 8 + n), gen.RESUELTO), (d(9, 3, 8 + n), gen.CERRADO)))
    tickets.append(ticket(40, d(8, 24), "Tecnico 05", "Baja", (d(8, 25), gen.RESUELTO), (d(8, 26), gen.CERRADO)))
    importar_diario(bd, coordinador, tickets, d(8, 24), d(9, 4, 23), tmp_path)
    return bd


def test_ca19_muestra_con_todos_los_p1_y_uno_por_tecnico(semana_cerrada, coordinador):
    semana = periodos.semana(2026, 36)
    muestra = met.muestra_de_semana(semana_cerrada, coordinador, semana, ahora=d(9, 7))
    motivos = dict(zip(muestra["ID"], muestra["Motivo"]))
    assert motivos[1] == "P1"
    assert motivos[2].startswith("Escalado más de")
    assert set(muestra["Técnico"]) >= {"Tecnico 01", "Tecnico 02", "Tecnico 03", "Tecnico 04", "Tecnico 05"}
    assert motivos[40] == "Al menos uno por técnico cada 2 semanas"  # cerrado la semana anterior
    assert len(muestra) >= 10
    # La muestra queda guardada: pedirla otra vez da la misma
    assert list(met.muestra_de_semana(semana_cerrada, coordinador, semana, ahora=d(9, 8))["ID"]) == list(muestra["ID"])


def test_agregar_y_quitar_de_la_muestra(semana_cerrada, coordinador):
    semana = periodos.semana(2026, 36)
    met.muestra_de_semana(semana_cerrada, coordinador, semana, ahora=d(9, 7))
    met.quitar_de_muestra(semana_cerrada, coordinador, semana, 1)
    met.agregar_a_muestra(semana_cerrada, coordinador, semana, 40)
    ids = set(met.muestra_de_semana(semana_cerrada, coordinador, semana)["ID"])
    assert 1 not in ids and 40 in ids


def test_not04_los_lunes(semana_cerrada, coordinador):
    lunes = d(9, 7, 8)
    assert "NOT-04" in [n.codigo for n in noti.pendientes(semana_cerrada, coordinador, noti.AL_INICIAR, lunes)]
    martes = d(9, 8, 8)
    assert "NOT-04" not in [n.codigo for n in noti.pendientes(semana_cerrada, coordinador, noti.AL_INICIAR, martes)]


# --- CAL-05 (CA-20) ---

def test_ca20_historico_con_flechas_y_variacion(bd, coordinador, tmp_path):
    # Atendidos por semana: S33=2, S34=4, S35=1, S36=3, S37=6
    cantidades = {33: 2, 34: 4, 35: 1, 36: 3, 37: 6}
    tickets, numero = [], 1
    for semana, cantidad in cantidades.items():
        lunes = datetime.fromisocalendar(2026, semana, 1).replace(hour=9)
        for i in range(cantidad):
            tickets.append(ticket(numero, lunes + timedelta(hours=i), "Tecnico 01", "Baja",
                                  (lunes + timedelta(days=1, hours=i), gen.RESUELTO)))
            numero += 1
    importar_diario(bd, coordinador, tickets, d(8, 10), d(9, 12, 23), tmp_path)
    tabla = met.historico_semanal(bd, coordinador, tecnico_id(bd, "Tecnico 01"), semanas=5, ahora=d(9, 13))
    fila = tabla.set_index("Semana")
    assert list(fila["Atendidos"]) == [2, 4, 1, 3, 6]
    assert fila.loc["2026-W34", "Atendidos vs. semana anterior"] == "▲ 100 %"
    assert fila.loc["2026-W35", "Atendidos vs. semana anterior"] == "▼ 75 %"
    assert fila.loc["2026-W37", "Atendidos vs. semana anterior"] == "▲ 100 %"
    # Promedio de las 4 semanas anteriores a la 37: (2 + 4 + 1 + 3) / 4 = 2,5 → 6 es +140 %
    assert fila.loc["2026-W37", "Atendidos vs. promedio 4 semanas"] == "▲ 140 %"
    assert fila.loc["2026-W33", "Atendidos vs. semana anterior"] == "—"


def test_historico_de_otro_tecnico_prohibido(bd, coordinador, tmp_path):
    importar_diario(bd, coordinador, [ticket(1, d(9, 1), "Tecnico 01", "Baja"), ticket(2, d(9, 1), "Tecnico 02", "Baja")],
                    d(9, 1), d(9, 1, 23), tmp_path)
    usuario = seguridad.crear_usuario(bd, coordinador, nombre="T2", perfil=seguridad.CONSULTA, pin="2222",
                                      tecnico_id=tecnico_id(bd, "Tecnico 02"))
    sesion = seguridad.iniciar_sesion(bd, usuario, "2222")
    with pytest.raises(ErrorPermiso):
        met.historico_semanal(bd, sesion, tecnico_id(bd, "Tecnico 01"))


# --- CAL-06 (CA-22) ---

def test_ca22_velocidad_normalizada_por_categoria(bd, coordinador, tmp_path):
    """Tecnico 01 atiende casos lentos (REC, ~40 h) y Tecnico 02 casos rápidos (APL, ~2 h).
    Ambos resuelven igual que el equipo en su categoría: índice 1, aunque sus horas difieran 20 veces."""
    tickets = []
    for n, horas in enumerate((36, 40, 44)):
        tickets.append(ticket(100 + n, d(9, 1, 8), "Tecnico 01", "Baja", (d(9, 1, 8) + timedelta(hours=horas), gen.RESUELTO)))
        tickets.append(ticket(200 + n, d(9, 1, 8), "Tecnico 03", "Baja", (d(9, 1, 8) + timedelta(hours=horas), gen.RESUELTO)))
    for n, horas in enumerate((1, 2, 3)):
        tickets.append(ticket(300 + n, d(9, 5, 8), "Tecnico 02", "Baja", (d(9, 5, 8) + timedelta(hours=horas), gen.RESUELTO)))
        tickets.append(ticket(400 + n, d(9, 5, 8), "Tecnico 03", "Baja", (d(9, 5, 8) + timedelta(hours=horas), gen.RESUELTO)))
    importar_diario(bd, coordinador, tickets, d(9, 1), d(9, 8, 23), tmp_path)
    cl.clasificar(bd, coordinador, [100, 101, 102, 200, 201, 202], {cl.CATEGORIA: "REC-01"})
    cl.clasificar(bd, coordinador, [300, 301, 302, 400, 401, 402], {cl.CATEGORIA: "APL-01"})
    indices = met.indices_velocidad(bd, periodos.mes(2026, 9))
    lento, rapido = tecnico_id(bd, "Tecnico 01"), tecnico_id(bd, "Tecnico 02")
    assert indices[lento] == indices[rapido] == 1.0
    tabla = met.rendimiento(bd, coordinador, periodos.mes(2026, 9), ahora=d(9, 30)).set_index("Técnico")
    assert tabla.loc["Tecnico 01", "Solución (h) ≈"] == 40.0 and tabla.loc["Tecnico 02", "Solución (h) ≈"] == 2.0
    assert tabla.loc["Tecnico 01", "Índice de velocidad"] == tabla.loc["Tecnico 02", "Índice de velocidad"]


# --- CAL-07 (CA-21) ---

@pytest.fixture
def equipo_evaluado(bd, coordinador, tmp_path):
    tickets = []
    for n in range(12):
        tickets.append(ticket(100 + n, d(9, 1, 8 + n % 8), "Tecnico 01", "Baja", (d(9, 2, 8 + n % 8), gen.CERRADO)))
        tickets.append(ticket(200 + n, d(9, 1, 8 + n % 8), "Tecnico 02", "Baja", (d(9, 3, 8 + n % 8), gen.CERRADO)))
    for n in range(4):  # Tecnico 03: menos de 10 tickets
        tickets.append(ticket(300 + n, d(9, 1, 9), "Tecnico 03", "Baja", (d(9, 2, 9), gen.CERRADO)))
    importar_diario(bd, coordinador, tickets, d(9, 1), d(9, 3, 23), tmp_path)
    for ticket_id in (100, 101, 102):
        evaluar_todo(bd, coordinador, ticket_id, ahora=d(9, 10))
    evaluar_todo(bd, coordinador, 200, ahora=d(9, 10), fallar=("G-02",))  # G-02 es crítico
    return bd


def test_ca21_ranking_excluye_muestra_pequena_y_solo_coordinador(equipo_evaluado, coordinador, consulta):
    resultado = met.ranking(equipo_evaluado, coordinador, periodos.mes(2026, 9), ahora=d(9, 30))
    nombres = [f.nombre for f in resultado.filas]
    assert nombres == ["Tecnico 01", "Tecnico 02"]
    assert ("Tecnico 03", "Muestra pequeña: 4 tickets (mínimo 10)") in resultado.excluidos
    uno, dos = resultado.filas
    assert uno.calidad == 100.0 and uno.evaluaciones == 3
    assert dos.calidad < 100 and dos.indice < uno.indice  # crítico fallido: −10 puntos
    assert resultado.aviso == met.AVISO_RANKING and resultado.promedio["indice"] is not None
    with pytest.raises(ErrorPermiso):
        met.ranking(equipo_evaluado, consulta, periodos.mes(2026, 9))


def test_ranking_respeta_incluir_en_ranking(equipo_evaluado, coordinador):
    with equipo_evaluado:
        equipo_evaluado.execute("UPDATE tecnico SET incluir_en_ranking = 0 WHERE nombre_glpi = 'Tecnico 02'")
    resultado = met.ranking(equipo_evaluado, coordinador, periodos.mes(2026, 9), ahora=d(9, 30))
    assert [f.nombre for f in resultado.filas] == ["Tecnico 01"]
    assert ("Tecnico 02", "Marcado «no incluir en ranking»") in resultado.excluidos


def test_tendencia_de_tres_meses(equipo_evaluado, coordinador):
    tabla = met.tendencia_ranking(equipo_evaluado, coordinador, periodos.mes(2026, 9))
    assert list(tabla.columns) == ["2026-07", "2026-08", "2026-09"]
    assert tabla.loc["Tecnico 01", "2026-09"] is not None


# --- CAL-08 ---

def test_incumplimiento_por_criterio(equipo_evaluado, coordinador):
    tabla = met.incumplimiento_por_criterio(equipo_evaluado, coordinador, periodos.mes(2026, 9)).set_index("Criterio")
    assert tabla.loc["G-02", "% incumplimiento equipo"] == 25.0  # 1 de 4 evaluaciones
    assert tabla.loc["G-02", "Crítico"] == "CRÍTICO"
    assert tabla.loc["G-02", "Capacitación"] == ""
    assert tabla.loc["G-02", "Tecnico 02 %"] == 100.0 and tabla.loc["G-02", "Tecnico 01 %"] == 0.0


# --- REP-06, REP-07, REP-10 (calidad) y resúmenes semanales ---

def solicitud(bd, sesion, carpeta, periodo, **extra):
    from core.reportes import catalogo
    return catalogo.SolicitudReporte(conexion=bd, sesion=sesion, periodo=periodo, carpeta_exportaciones=carpeta,
                                     ahora=d(9, 30), **extra)


def test_rep06_historico_por_tecnico(equipo_evaluado, coordinador, consulta, tmp_path):
    from openpyxl import load_workbook
    from pypdf import PdfReader
    from core.analisis.filtros import Filtros
    from core.reportes import catalogo
    septiembre = periodos.mes(2026, 9)
    excel = catalogo.generar(solicitud(equipo_evaluado, coordinador, tmp_path, septiembre), "REP-06", catalogo.EXCEL)
    assert {"Tecnico 01", "Tecnico 02", "Tecnico 03"} <= set(load_workbook(excel).sheetnames)
    uno = catalogo.generar(solicitud(equipo_evaluado, coordinador, tmp_path, septiembre,
                                     filtros=Filtros(tecnico_id=tecnico_id(equipo_evaluado, "Tecnico 01"))),
                           "REP-06", catalogo.PDF)
    texto = "\n".join(pag.extract_text() for pag in PdfReader(uno).pages)
    assert "REP-06 Calidad de soporte" in texto and "Tecnico 01" in texto and "Tecnico 02" not in texto
    with pytest.raises(ErrorPermiso):
        catalogo.generar(solicitud(equipo_evaluado, consulta, tmp_path, septiembre), "REP-06", catalogo.PDF)


def test_rep07_comparativo_solo_coordinador(equipo_evaluado, coordinador, consulta, tmp_path):
    from openpyxl import load_workbook
    from pypdf import PdfReader
    from core.reportes import catalogo
    septiembre = periodos.mes(2026, 9)
    excel = catalogo.generar(solicitud(equipo_evaluado, coordinador, tmp_path, septiembre), "REP-07", catalogo.EXCEL)
    libro = load_workbook(excel)
    assert libro.sheetnames == ["Ranking", "Incumplimiento"]
    valores = [c for fila in libro["Ranking"].iter_rows(values_only=True) for c in fila if c is not None]
    assert "Promedio del equipo" in valores and "Tecnico 03" in valores  # excluido, con su motivo
    pdf = catalogo.generar(solicitud(equipo_evaluado, coordinador, tmp_path, septiembre), "REP-07", catalogo.PDF)
    texto = " ".join(pag.extract_text().replace("\n", " ") for pag in PdfReader(pdf).pages)
    assert met.AVISO_RANKING in texto
    assert "REP-07" not in {r.codigo for r in catalogo.disponibles(consulta)}
    with pytest.raises(ErrorPermiso):
        catalogo.generar(solicitud(equipo_evaluado, consulta, tmp_path, septiembre), "REP-07", catalogo.PDF)


def test_rep10_incluye_la_calidad_del_area(equipo_evaluado, coordinador, tmp_path):
    from openpyxl import load_workbook
    from core.reportes import paquete
    resultado = paquete.generar(equipo_evaluado, coordinador, periodos.mes(2026, 9), tmp_path, ahora=d(10, 1))
    hoja = load_workbook(resultado.excel)["Calidad del área"]
    valores = [c for fila in hoja.iter_rows(values_only=True) for c in fila if c is not None]
    assert "Evaluaciones de calidad del mes" in valores and "Criterios con más incumplimiento" in valores
    assert any(isinstance(v, str) and v.startswith("KPI-14") for v in valores)
    assert any(isinstance(v, str) and v.startswith("KPI-15") for v in valores)
    assert cal.NOMBRE_RESULTADO[cal.CONFORME] in valores
    assert not any(isinstance(v, str) and v.startswith("Tecnico") for v in valores)  # sin datos por técnico


def test_resumen_semanal_solo_con_datos_propios(equipo_evaluado, coordinador, consulta, tmp_path):
    from email import policy
    from email.parser import BytesParser
    from core.reportes import paquete
    evaluar_todo(equipo_evaluado, coordinador, 201, ahora=d(9, 10), retroalimentacion="Documentar la causa raíz.")
    semana = periodos.semana(2026, 37)  # 7 al 13 de septiembre: solo hay evaluaciones
    assert paquete.semana_resumen(d(9, 16)) == semana
    rutas = paquete.resumenes_semanales(equipo_evaluado, coordinador, tmp_path, ahora=d(9, 16))
    assert sorted(r.name for r in rutas) == ["resumen_Tecnico_01.eml", "resumen_Tecnico_02.eml"]
    mensaje = BytesParser(policy=policy.default).parsebytes(
        next(r for r in rutas if r.name == "resumen_Tecnico_02.eml").read_bytes())
    assert mensaje["X-Unsent"] == "1" and not mensaje["To"]
    cuerpo = mensaje.get_body(preferencelist=("plain",)).get_content()
    assert "Hola, Tecnico 02" in cuerpo and "Ticket 200" in cuerpo and "Ticket 201" in cuerpo
    assert "Documentar la causa raíz." in cuerpo
    assert "Tecnico 01" not in cuerpo and "ranking" not in cuerpo.lower() and "Índice" not in cuerpo
    semana36 = paquete.resumenes_semanales(equipo_evaluado, coordinador, tmp_path, periodos.semana(2026, 36), d(9, 16))
    assert len(semana36) == 3  # los tres técnicos atendieron tickets
    with pytest.raises(ErrorPermiso):
        paquete.resumenes_semanales(equipo_evaluado, consulta, tmp_path, ahora=d(9, 16))


def test_rep11_trae_la_evaluacion_vigente(equipo_evaluado, coordinador, tmp_path):
    import pandas as pd
    from core.reportes import catalogo
    ruta = catalogo.generar(solicitud(equipo_evaluado, coordinador, tmp_path, periodos.mes(2026, 9)), "REP-11", catalogo.CSV)
    tabla = pd.read_csv(ruta, sep=";", encoding="utf-8-sig").set_index("ID")
    assert tabla.loc[100, "Evaluación %"] == 100.0
    assert tabla.loc[200, "Críticos fallidos"] == 1
    assert pd.isna(tabla.loc[300, "Evaluación %"])
